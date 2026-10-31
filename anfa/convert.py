"""Part 3: from parser output to structured documents you can trust.

A parser tells you what it found on each page. Before anything is indexed, this step turns that
into the document a reader would recognize, and checks the result:

1. Drop page furniture: running headers and footers, and the Federal Register's print-shop lines.
2. Keep only the document itself: a directive printed in the Federal Register shares its first
   and last pages with the neighbouring rules.
3. Apply corrections a person made after checking the page, and list the scanned pages nobody
   has checked yet. The OCR's own confidence can't be used to triage: it scored every page of
   the scanned work orders 0.99, misread signatures included.
4. Repair what the parser broke in a predictable way: AD numbers split at a line break, a
   section heading read as a paragraph.
5. Rejoin tables split across pages, and attach their caption, or the sheet they came from.
6. Put back what the parser left out, from the source file: slide notes.
7. Describe figures with a vision model, so what they show becomes text.
8. Report what was changed and what still needs a person.

The structured document keeps each element's type. Part 3 still renders it to text for the
size-based chunker; Part 4 uses the types directly.
"""

from __future__ import annotations

import hashlib
import json
import re
from dataclasses import dataclass, field
from pathlib import Path

from anfa.corpus import CORPUS_DIR, ParsedDocument
from anfa.providers import Describer

# The Federal Register's print-shop lines and running heads, which the parser reads as body text.
PRINT_LINES = re.compile(
    r"^(VerDate |BILLING CODE |\[FR Doc\. |Federal Register\s*/\s*Vol\.|\d{1,5}$|\S+ on \w*PROD\w* with \w+$)"
)
# An AD number whose hyphen was taken for a hyphenated line break: "AD 200324-13", "AD 2003-2413".
SPLIT_AD_NUMBER = re.compile(r"(?<=\bAD )((?:19|20)\d{2})-?(\d{2})-?(\d{2})\b")
# A TCDS section starts with a roman numeral and a model: "IV. Model 172D, ...".
TCDS_SECTION = re.compile(r"^[IVXL]+\. Models? ")
CAPTION = re.compile(r"^TABLE [\w-]+\.", re.IGNORECASE)


@dataclass
class Element:
    kind: str  # paragraph, table, figure, slide
    page: int | None
    data: dict

    def render(self) -> str:
        """Text for search and for the model. Part 4 refines what each kind supplies."""
        d = self.data
        if self.kind == "table":
            lines = [d["scope"]] if d.get("scope") else []
            if d["header"]:
                lines.append("| " + " | ".join(d["header"]) + " |")
                lines.append("|" + " --- |" * len(d["header"]))
            lines += ["| " + " | ".join(row) + " |" for row in d["rows"]]
            if d.get("note"):
                lines.append(d["note"])
            return "\n".join(lines)
        if self.kind == "figure":
            return "\n".join(t for t in (d["caption"], f"Figure description: {d['description']}") if t)
        if self.kind == "slide":
            notes = f"\nSpeaker notes: {d['notes']}" if d["notes"] else ""
            return f"{d['title']}\n{d['text']}{notes}"
        return d["text"]


@dataclass
class Section:
    number: str
    title: str
    elements: list[Element] = field(default_factory=list)
    # What the page is, when one file holds several documents (a reviewer's page title).
    context: str = ""
    # Part 6: what the section applies to, when that differs from the document (a page's aircraft).
    meta: dict = field(default_factory=dict)


@dataclass
class IngestReport:
    source_id: str
    file: str
    removed_furniture: int = 0
    removed_neighbours: int = 0
    merged_tables: int = 0
    repaired: list[str] = field(default_factory=list)
    restored: list[str] = field(default_factory=list)
    corrected: list[str] = field(default_factory=list)
    needs_review: list[str] = field(default_factory=list)
    described_figures: int = 0

    def lines(self) -> list[str]:
        notes = []
        if self.removed_furniture:
            notes.append(f"removed {self.removed_furniture} page headers, footers, and print-shop lines")
        if self.removed_neighbours:
            notes.append(f"removed {self.removed_neighbours} elements of the neighbouring rules on the same pages")
        if self.merged_tables:
            notes.append(f"rejoined {self.merged_tables} table part(s) split across pages")
        if self.described_figures:
            notes.append(f"described {self.described_figures} figure(s)")
        notes += [f"repaired: {r}" for r in self.repaired]
        notes += [f"restored from the source file: {r}" for r in self.restored]
        notes += [f"applied reviewed correction: {c}" for c in self.corrected]
        notes += [f"NEEDS REVIEW: {n}" for n in self.needs_review]
        return notes or ["no changes"]


@dataclass
class StructuredDocument:
    source_id: str
    meta: dict[str, str]
    sections: list[Section]
    report: IngestReport


# --- Reviewed corrections -------------------------------------------------------------------------

def fingerprint(element: dict) -> str:
    """What a correction was written against. If the parse changes, the correction no longer applies."""
    content = {k: v for k, v in element.items() if k != "id"}
    return hashlib.sha256(json.dumps(content, sort_keys=True, ensure_ascii=False).encode()).hexdigest()[:12]


def load_corrections(directory: Path = CORPUS_DIR) -> dict:
    path = directory / "corrections.json"
    return json.loads(path.read_text(encoding="utf-8")) if path.exists() else {}


# --- Steps that look at the whole document ---------------------------------------------------------

def own_rule(doc: ParsedDocument) -> tuple[int, int] | None:
    """The span of a Federal Register printing that is this directive, as (first, last) indexes.

    A rule starts with "DEPARTMENT OF TRANSPORTATION" just above its docket line, which names the
    directive, and ends with its "[FR Doc. number ...]" line. Where the parser read two rules'
    columns interleaved, part of the other rule stays inside the span: the reading order gives no
    way to separate them.
    """
    elements = doc.elements
    ad = doc.meta.get("ad_number")
    number = re.search(r"document ([\w-]+)", doc.meta.get("federal_register", ""))
    if not ad:
        return None
    named = re.compile(rf"\bAD {re.escape(ad)}\b")
    docket = next((i for i, e in enumerate(elements)
                   if "Amendment 39" in e.get("text", "") and named.search(e["text"])), None)
    if docket is None:
        return None
    first = next((i for i in range(docket, max(-1, docket - 8), -1)
                  if elements[i].get("text", "").strip().upper() == "DEPARTMENT OF TRANSPORTATION"), docket)
    # Without its "[FR Doc." line in the parse, a rule ends where the next one starts.
    next_rule = next((i - 1 for i in range(docket, len(elements))
                      if elements[i].get("text", "").strip().upper() == "DEPARTMENT OF TRANSPORTATION"), len(elements) - 1)
    last = next((i for i in range(docket, len(elements))
                 if number and elements[i].get("text", "").startswith(f"[FR Doc. {number.group(1)} ")), next_rule)
    return first, last


def repair_ad_numbers(text: str) -> tuple[str, int]:
    """Put back the hyphens of AD numbers. Text quoted between '' marks is left alone: a correction
    notice quotes the misprint it corrects."""
    count = 0

    def fix(match: re.Match) -> str:
        nonlocal count
        if text[:match.start()].count("''") % 2:  # inside a quotation
            return match.group(0)
        if re.search(r"(European Aviation Safety Agency|EASA) $", text[:match.start() - 3]):
            return match.group(0)  # EASA numbers its directives 2007-0232: nothing is broken
        fixed = "-".join(match.groups())
        count += fixed != match.group(0)
        return fixed

    return SPLIT_AD_NUMBER.sub(fix, text), count


def slide_notes(path: Path) -> dict[int, str]:
    from pptx import Presentation

    return {n: s.notes_slide.notes_text_frame.text.strip() for n, s in enumerate(Presentation(path).slides, start=1)
            if s.has_notes_slide}


def sheet_names(path: Path) -> list[tuple[str, list[str]]]:
    """Each sheet's name and its first row, so a parsed table can be matched to the sheet it came from."""
    from openpyxl import load_workbook

    book = load_workbook(path, read_only=True)
    sheets = [(ws.title, [str(v) if v is not None else "" for v in next(ws.iter_rows(values_only=True), ())])
              for ws in book.worksheets]
    book.close()
    return sheets


# --- Conversion -----------------------------------------------------------------------------------

def convert(doc: ParsedDocument, describer: Describer | None, corrections: dict) -> StructuredDocument:
    report = IngestReport(doc.source_id, doc.file)
    fixes = corrections.get("corrections", {}).get(doc.source_id, {})
    checked_pages = set(corrections.get("reviewed_pages", {}).get(doc.source_id, []))
    page_titles = corrections.get("page_titles", {}).get(doc.source_id, {})
    page_meta = corrections.get("page_meta", {}).get(doc.source_id, {})
    source = CORPUS_DIR / "source" / doc.file
    notes = slide_notes(source) if doc.format == "pptx" else {}
    sheets = sheet_names(source) if doc.format == "xlsx" else []
    sections: list[Section] = []
    pending_caption = ""

    elements = doc.elements
    # 2. A directive keeps only its own rule.
    if span := own_rule(doc):
        first, last = span
        outside = elements[:first] + elements[last + 1:]
        report.removed_neighbours = sum(not _furniture(e) for e in outside)
        report.removed_furniture += len(outside) - report.removed_neighbours
        elements = elements[first:last + 1]

    def new_section(number: str, title: str) -> Section:
        page = str(raw.get("page"))  # the element being placed
        return Section(number, title, context=page_titles.get(page, ""), meta=page_meta.get(page, {}))

    def section() -> Section:
        if not sections:
            sections.append(new_section("", ""))
        return sections[-1]

    read_by_ocr = ocr_pages(doc, elements)
    repaired_numbers = 0
    for raw in elements:
        kind = raw["type"]
        text = raw.get("text")

        # 1. Page furniture repeats on every page and carries no evidence.
        if _furniture(raw):
            report.removed_furniture += 1
            continue

        # 3. Reviewed corrections, applied only to the element they were written against.
        if fix := fixes.get(raw["id"]):
            if fix["was"] != fingerprint(raw):
                report.needs_review.append(f"{raw['id']}: the parse changed since the correction was written; check it again")
            else:
                report.corrected.append(f"{raw['id']} (page {raw.get('page')}): {fix['why']}")
                if fix.get("drop"):
                    continue
                if kind == "table":
                    raw = {**raw, **fix["table"]}
                else:
                    text = fix["text"]
                kind = fix.get("type", kind)

        # 4. Breaks the parser makes in a predictable way.
        if text and doc.meta.get("doc_type") == "airworthiness directive":
            text, n = repair_ad_numbers(text)
            repaired_numbers += n
        if kind == "paragraph" and doc.meta.get("doc_type") == "type certificate data sheet" and TCDS_SECTION.match(text):
            kind = "heading"
            report.repaired.append(f"{raw['id']} (page {raw.get('page')}) is a section heading: {text[:40]}...")

        if kind == "heading":
            number, title = _heading(text)
            sections.append(new_section(number, title))
            continue
        if kind == "slide":
            # 6. The parser keeps a slide's title and text and drops its notes; the source file has them.
            slide_notes_text = raw["notes"] or notes.get(raw["slide"], "")
            if slide_notes_text and not raw["notes"]:
                report.restored.append(f"speaker notes of slide {raw['slide']}")
            sections.append(Section(str(raw["slide"]), raw["title"]))
            section().elements.append(Element("slide", None, {
                "slide": raw["slide"], "title": raw["title"], "text": raw["text"], "notes": slide_notes_text,
            }))
            continue

        current = section()
        previous = current.elements[-1] if current.elements else None
        if kind == "paragraph":
            # 5. A caption next to a table belongs to the table.
            if CAPTION.match(text):
                if previous and previous.kind == "table" and not previous.data["scope"]:
                    previous.data["scope"] = text
                else:
                    pending_caption = text
                continue
            current.elements.append(Element("paragraph", raw.get("page"), {"text": text}))
        elif kind == "table":
            # 5. A table that repeats the previous table's header, or has none, continues it on the next page.
            if previous and previous.kind == "table" and raw["header"] in (None, previous.data["header"]):
                previous.data["rows"].extend(raw["rows"])
                report.merged_tables += 1
                continue
            scope = raw.get("scope") or pending_caption
            if sheets:
                name = next((title for title, first_row in sheets if first_row[:len(raw["header"] or [])] == raw["header"]), "")
                scope = f"Sheet: {name}" if name else scope
            current.elements.append(Element("table", raw.get("page"), {
                "header": raw["header"], "rows": [list(r) for r in raw["rows"]], "scope": scope, "note": raw.get("note", ""),
            }))
            pending_caption = ""
        elif kind == "figure":
            # 7. A figure's content is in the image; a vision model turns it into text.
            description = describer.describe(CORPUS_DIR / raw["image"], raw["caption"], doc.meta["title"]) if describer else ""
            if description is None:
                report.needs_review.append(f"{raw['id']} (page {raw.get('page')}): no usable figure description")
                description = ""
            report.described_figures += bool(description)
            current.elements.append(Element("figure", raw.get("page"), {
                "caption": raw["caption"], "image": raw["image"], "description": description,
            }))

    if repaired_numbers:
        report.repaired.append(f"{repaired_numbers} AD number(s) split at a line break")
    if sheets:
        report.restored.append("sheet names, as table scope")
    # 3. A scanned page is trusted once a person has checked it against the image.
    unchecked = [p for p in read_by_ocr if p not in checked_pages]
    if unchecked:
        report.needs_review.append(f"OCR text not yet checked against the scan, pages {_ranges(unchecked)}")
    return StructuredDocument(doc.source_id, doc.meta, sections, report)


def _furniture(element: dict) -> bool:
    return element["type"] in ("page_header", "page_footer") or bool(PRINT_LINES.search(element.get("text") or ""))


def ocr_pages(doc: ParsedDocument, elements: list[dict]) -> list[int]:
    """Pages whose text the parser read from pixels.

    The parser runs OCR on any bitmap, so an OCR score on a page doesn't mean its text came from
    OCR. A page did when the parser found much more text on it than the PDF's own text layer holds,
    or when the whole document was OCR'd on purpose. A scan parsed with OCR off was read from its
    own text layer, which an earlier OCR wrote: all of it.
    """
    if doc.ocr == "off":
        return sorted({e["page"] for e in elements if e.get("page")})
    scored = sorted({e["page"] for e in elements if "ocr_confidence" in e and e.get("page")})
    if doc.ocr == "full-page" or not scored:
        return scored
    import pypdfium2

    pdf = pypdfium2.PdfDocument(CORPUS_DIR / "source" / doc.file)
    parsed = {p: sum(len(e.get("text") or "") for e in elements if e.get("page") == p) for p in scored}
    layer = {p: len(pdf[p - 1].get_textpage().get_text_range().strip()) for p in scored}
    pdf.close()
    return [p for p in scored if layer[p] < parsed[p] / 2]


def _heading(text: str) -> tuple[str, str]:
    """Split a numbered heading ("7-44. CLOSE-TOLERANCE BOLTS.", "IV. Model 172D") into number and title."""
    number, _, title = text.partition(" ")
    if re.fullmatch(r"[\w().-]*\d[\w().-]*\.?|[IVXL]+\.|§", number):
        return number.rstrip("."), title
    return "", text


def _ranges(pages: list[int]) -> str:
    runs: list[list[int]] = []
    for p in pages:
        if runs and p == runs[-1][-1] + 1:
            runs[-1].append(p)
        else:
            runs.append([p])
    return ", ".join(str(r[0]) if len(r) == 1 else f"{r[0]}-{r[-1]}" for r in runs)
