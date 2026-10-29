"""Build Anfa Aero Services' own documents: the fictional company layer on top of the FAA corpus.

Anfa Aero Services is a fictional maintenance shop that looks after the Cessna 172 fleet of a
fictional flight school, Tadla Flight Academy. The FAA documents in corpus/source/faa/ are real;
these are not. They are written to fit the real ones: the serial numbers, part numbers, and
equipment fall inside, outside, and just next to the ranges in the real airworthiness directives.

    FLEET-2026-10   Excel workbook: aircraft, installed equipment, time-controlled parts
    PROC-R7         Word: the shop's procedures manual (excerpt)
    TRAIN-AD        slides whose key information is in the speaker notes
    WO-SCAN         scanned work orders: page images only, no text layer
    PRICES-2026     Excel workbook, readable by management only
    WO-yyyy-nnnn    work orders printed from the maintenance system, one PDF each, 2024 to 2026
    LOG-CN-TFx      each aircraft's logbook entries for the same period, printed as a table
    ADREC-CN-TFx    each aircraft's AD compliance record, Excel

The work orders, logbooks, and AD compliance records come from anfa_history.py, which derives them
from the fleet's hours; build_eval.py computes the evaluation answers from the same module.

Metadata (access, revision, dates) is not inside the files, as in a real document management
system: it goes to corpus/source/catalog.json, next to the FAA entries.

    uv run --group corpus python scripts/build_anfa.py
"""

from __future__ import annotations

import io
import json
import random
import zipfile
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / "corpus" / "source"
OUT = SOURCE / "anfa"
FIXED = datetime(2026, 10, 26, 9, 0, 0)
COMPANY = "Anfa Aero Services"

# --- The fleet ------------------------------------------------------------------------------
# Hours as of 26 October 2026. Each line notes, for the reader of this script, which real ADs
# the aircraft falls inside or just outside of. None of these notes go into the documents.

FLEET = [
    # CN-TFA: inside AD 2013-11-11 (172S8001-172S11256); three short of AD 2012-02-02 (from 172S9491).
    ("CN-TFA", "172S", "172S9488", 2003, 10140.6, "Lycoming IO-360-L2A", "L-31456-51A", 1640.6, "Casablanca",
     "Cessna SB04-28-03 and fuel return kit MK172-28-01 incorporated in 2005."),
    # CN-TFB: inside AD 2013-11-11, AD 2012-02-02, AD 2007-08-03, AD 2008-05-09.
    ("CN-TFB", "172S", "172S10237", 2006, 8455.1, "Lycoming IO-360-L2A", "L-33872-51A", 455.1, "Casablanca", ""),
    # CN-TFC: same model, outside AD 2013-11-11 (above 172S11256) and AD 2013-03-15 (above 172S11193).
    ("CN-TFC", "172S", "172S11290", 2013, 6210.3, "Lycoming IO-360-L2A", "L-37115-51A", 2210.3, "Beni Mellal", ""),
    # CN-TFD: inside AD 2013-11-11 and AD 2012-02-02; fourteen above AD 2012-22-01 (to 17281187).
    ("CN-TFD", "172R", "17281201", 2003, 11870.4, "Lycoming IO-360-L2A", "L-31377-51A", 1370.4, "Casablanca", ""),
    # CN-TFE: inside AD 2013-11-11, AD 2007-05-10, and the KAP 140 ADs; its Aspen display meets AD 2020-16-08.
    ("CN-TFE", "172R", "17280744", 1999, 13520.8, "Lycoming IO-360-L2A", "L-28940-51A", 1120.8, "Casablanca", ""),
    # CN-TFF: built by Textron Aviation (TCDS 3A12 note 5); its AWI muffler is inside AD 2018-02-04.
    ("CN-TFF", "172S", "172S11845", 2017, 3980.0, "Lycoming IO-360-L2A", "L-39552-51A", 3980.0, "Beni Mellal", ""),
    # CN-TFG: a 172N converted to a diesel engine under STC SA01303WI: AD 2011-06-02 applies, the
    # TAE 125-01 ADs (2009-24-10, 2010-04-06) do not, because its engine is a TAE 125-02-99.
    ("CN-TFG", "172N", "17272851", 1978, 14960.2, "Technify (Thielert) TAE 125-02-99", "02-02-1173", 610.2, "Beni Mellal",
     "Converted from Lycoming O-320-H2AD under STC SA01303WI in 2011."),
]

EQUIPMENT = [
    ("CN-TFA", "Transponder", "Garmin", "GTX 330", "77412093", "3.06", "2003", "Software upgraded from 3.04 in March 2005."),
    # The AWI mufflers inside AD 2018-02-04 were replaced in 2018 (anfa_history.py, COMPLIANCE). CN-TFA's row
    # was updated; CN-TFF's was not, and still shows the affected muffler (discrepancy D2).
    ("CN-TFA", "Muffler", "Aerospace Welding Minneapolis (AWI)", "A1754001-25", "39274", "", "2018-03-12",
     "Replaced S/N 34640 under AD 2018-02-04."),
    ("CN-TFB", "Integrated flight deck", "Garmin", "G1000", "", "", "2006", "Factory installed."),
    ("CN-TFC", "Integrated flight deck", "Garmin", "G1000", "", "", "2013", "Factory installed."),
    ("CN-TFC", "Muffler", "Aerospace Welding Minneapolis (AWI)", "A1754001-23", "34729", "", "2019-09-02", ""),
    ("CN-TFD", "Transponder", "Garmin", "GTX 330", "77418811", "3.06", "2003", ""),
    ("CN-TFD", "Autopilot", "Bendix/King", "KAP 140, P/N 065-00176-7702", "", "", "2003", "Factory installed."),
    ("CN-TFE", "Autopilot", "Bendix/King", "KAP 140, P/N 065-00176-5402", "", "SW MOD 03/01", "1999",
     "Factory installed. Software updated under AD 2004-15-18 in 2004 (M decal)."),
    ("CN-TFE", "Primary flight display", "Aspen Avionics", "EFD1000 PFD, P/N 910-00001-011", "EP0141287", "2.10.1", "2019-05-20",
     "Standby airspeed indicator and altimeter retained; standby attitude indicator removed at installation."),
    ("CN-TFF", "Integrated flight deck", "Garmin", "G1000 NXi", "", "", "2017", "Factory installed."),
    ("CN-TFF", "Muffler", "Aerospace Welding Minneapolis (AWI)", "A1754001-25", "38190", "", "2017", "Factory installed."),
    ("CN-TFG", "Engine", "Technify (Thielert)", "TAE 125-02-99", "02-02-1173", "", "2011-03-15", "STC SA01303WI."),
    ("CN-TFG", "Transponder", "Garmin", "GTX 330", "77420156", "3.06", "2011-03-15", ""),
]

# Engine oil pressure switch P/N 83278: 3,000 hours under AD 2013-11-11, and under the shop's own
# policy (procedures manual 3.2) for the two 172S outside the AD's serial range.
TIME_CONTROLLED = [
    ("CN-TFA", "Engine oil pressure switch", "83278", "", "2024-02-14", 7210.4, "3,000 h", "AD 2013-11-11"),
    ("CN-TFB", "Engine oil pressure switch", "83278", "", "2023-05-09", 5880.0, "3,000 h", "AD 2013-11-11"),
    ("CN-TFC", "Engine oil pressure switch", "83278", "", "2022-11-21", 3350.7, "3,000 h", "Procedures manual 3.2"),
    ("CN-TFD", "Engine oil pressure switch", "83278", "", "2025-03-03", 9940.2, "3,000 h", "AD 2013-11-11"),
    # The work order (WO-2021-0544) says 10,650.2 hours. The register copied it with two digits swapped.
    ("CN-TFE", "Engine oil pressure switch", "83278", "", "2021-08-30", 10560.2, "3,000 h", "AD 2013-11-11"),
    ("CN-TFF", "Engine oil pressure switch", "83278", "", "2022-07-18", 2410.5, "3,000 h", "Procedures manual 3.2"),
    ("CN-TFA", "ELT battery", "452-6505", "", "2021-01-19", 4755.3, "6 years", "ELT manufacturer"),
    ("CN-TFE", "ELT battery", "452-6505", "", "2020-11-25", 9810.0, "6 years", "ELT manufacturer"),
]

PRICES = [
    ("83278", "Engine oil pressure switch", "Textron Aviation", 1180, 1650, 12, "Keep two in stock (procedures manual 3.2)."),
    ("A1754001-25", "Muffler assembly", "Aerospace Welding Minneapolis", 21400, 26900, 35, ""),
    ("A1754001-23", "Muffler assembly", "Aerospace Welding Minneapolis", 20800, 26100, 35, ""),
    ("452-6505", "ELT battery pack", "ACK Technologies", 2350, 3100, 20, ""),
    ("CH48110-1", "Oil filter", "Champion Aerospace", 310, 420, 7, ""),
    ("REM37BY", "Spark plug", "Champion Aerospace", 290, 395, 7, "Sold per plug; eight per engine."),
    ("065-00176-7702", "KAP 140 autopilot computer (exchange)", "Honeywell (Bendix/King)", 61500, 74900, 45, "Exchange price; core returned."),
    ("910-00001-011", "Aspen EFD1000 PFD (exchange)", "Aspen Avionics", 58200, 71000, 30, "Exchange price; core returned."),
    ("LABOR-HR", "Labour, per hour", COMPANY, 0, 850, 0, "Flight school contract rate: 780 per hour."),
]

# --- Procedures manual ----------------------------------------------------------------------

PROCEDURES = [
    ("1 Purpose and scope", [
        "This manual describes how Anfa Aero Services maintains the aircraft it looks after. It covers the Cessna "
        "172 fleet of Tadla Flight Academy, based in Casablanca and Beni Mellal. This excerpt holds the sections "
        "that mechanics use every week; the full manual is kept by the quality manager.",
        "Where this manual and a manufacturer's instruction or an airworthiness directive differ, the directive "
        "comes first, then the manufacturer's instruction, then this manual. This manual may be stricter than "
        "either; it may never be less strict.",
    ]),
    ("2 Airworthiness directives", [
        "2.1 Weekly review. Every Monday the quality manager checks the FAA's list of new and revised "
        "airworthiness directives for the Cessna 172 series, the engines, propellers, and appliances installed "
        "in the fleet. The FAA is the state of design for every type we maintain, and its directives are "
        "applied as mandatory.",
        "2.2 Corrections. A directive can be corrected after publication, in a separate Federal Register notice. "
        "Always read a directive together with its corrections before deciding whether it applies.",
        "2.3 Applicability. A directive applies to an aircraft only when every condition in its applicability "
        "paragraph is met: model, serial number range, installed equipment, part number, software version, and "
        "any modification or service bulletin it names. Check each condition against the fleet register and the "
        "aircraft records, and write down which conditions you checked.",
        "2.4 Records. Each aircraft has an AD compliance record in the format suggested in AC 43-9D, appendix A: "
        "the directive number and revision, the method of compliance, the date and hours, and for recurring "
        "directives the next time due. A directive found not applicable is recorded too, with the reason.",
    ]),
    ("3 Life-limited and time-controlled parts", [
        "3.1 Tracking. Time-controlled parts are listed in the fleet register, sheet Time-controlled parts, with "
        "the airframe hours at installation. The hours used on a part are the airframe's current total time "
        "minus the airframe hours at installation.",
        "3.2 Engine oil pressure switch. AD 2013-11-11 limits the engine oil pressure switch, P/N 83278, to 3,000 "
        "hours time-in-service on the airplanes within its applicability. Anfa applies the same 3,000-hour limit to "
        "every Cessna 172R and 172S it maintains, including those outside the directive's serial number ranges. "
        "Plan the replacement when a switch reaches 2,900 hours. Keep two switches in stock.",
        "3.3 When the records disagree. If the fleet register and a work order or logbook give different hours for "
        "the same installation, use the lower airframe hours at installation until the quality manager has "
        "resolved the difference. Never use the figure that gives the part more time.",
    ]),
    ("4 Torque", [
        "4.1 Use the manufacturer's maintenance manual for the torque of any fastener it lists. When it gives no "
        "value, use AC 43.13-1B, chapter 7, Table 7-1, and read the notes printed with that table: they limit when "
        "its values apply.",
        "4.2 Record the torque wrench's identification number on the work order. Wrenches are calibrated every "
        "12 months; a wrench past its calibration date is not used.",
    ]),
    ("5 Maintenance records and return to service", [
        "5.1 Every maintenance record entry contains what 14 CFR 43.9 requires: a description of the work, the date "
        "it was completed, the name of the person who did it, and the signature and certificate number of the "
        "person approving the aircraft for return to service.",
        "5.2 Write airframe hours as read from the aircraft's hour meter, to one decimal place, and check the figure "
        "you copied from the work order into the register before signing.",
        "5.3 Work orders are scanned and kept with the aircraft records. The paper originals are kept for two years.",
    ]),
    ("6 Defect reporting", [
        "6.1 Any failure, malfunction, or defect that could affect safety is reported to the quality manager the same "
        "day. The quality manager reports it to the authority and, for FAA type-certificated products, through the "
        "FAA's Service Difficulty Reporting System.",
    ]),
    ("7 Receiving parts", [
        "7.1 A new or overhauled part is accepted only with an authorized release certificate, such as FAA Form "
        "8130-3. Check that the part number and serial number on the certificate match the part.",
    ]),
]

REVISIONS = [
    ("7", "2026-06-01", "Section 3.2 extended to every Cessna 172R and 172S; section 3.3 added."),
    ("6", "2025-09-29", "Section 2.4 and section 7 refer to AC 43-9D, which replaced AC 43-9C."),
    ("5", "2024-01-15", "Section 4.2: torque wrench calibration every 12 months instead of 24."),
]

# --- Training deck ---------------------------------------------------------------------------

SLIDES = [
    ("Airworthiness directives at Anfa",
     "New mechanics, first month\nQuality: Karim Bennani",
     "This deck is for mechanics in their first month. It covers how we decide whether a directive applies and how "
     "we record it. Ask the quality manager about anything it does not cover."),
    ("Where directives come from",
     "The state of design issues them: the FAA for our Cessnas\nPublished in the Federal Register\nRevisions and corrections follow",
     "Corrections are published as separate notices, sometimes weeks later. A correction can change the "
     "applicability or a compliance time. AD 2020-16-08 on the Aspen displays is an example: read the rule and "
     "its correction together."),
    ("Check applicability before you sign",
     "Model and serial number range\nInstalled equipment and part number\nSoftware version, modifications",
     "Two traps in our own fleet. CN-TFC is a 172S like CN-TFA and CN-TFB, but its serial number is above the range "
     "of AD 2013-11-11. CN-TFD has a KAP 140 autopilot, but its serial number is outside the range of the KAP 140 "
     "directives. Equipment alone is not enough, and serial numbers alone are not enough."),
    ("The engine oil pressure switch",
     "P/N 83278, 3,000 hours\nReplace with a new, zero-time switch\nRecord part number, date, and airframe hours",
     "The 50-hour allowance in paragraph (g)(6) of AD 2013-11-11 applied once, to switches that were already at or "
     "near the limit when the directive took effect on 1 August 2013. It is not a grace period you can use today. "
     "Our manual, section 3.2, applies the same 3,000-hour limit to every 172R and 172S we maintain."),
    ("Torque",
     "Maintenance manual first\nAC 43.13-1B Table 7-1 when it gives no value\nRecord the wrench number",
     "Read the notes printed with Table 7-1 before you use it: they limit when its values apply. If a value looks "
     "wrong, stop and ask; do not guess between two columns."),
    ("Records",
     "14 CFR 43.9 content in every entry\nAD compliance record for each aircraft\nHours to one decimal place",
     "Check every number you copy. A swapped pair of digits in an installation time moves the next due time by "
     "90 hours, and nobody notices until the records are audited."),
]

# --- Scanned work orders ----------------------------------------------------------------------

WORK_ORDERS = [
    ("WO-2021-0544", "30 August 2021", "CN-TFE", "Cessna 172R", "17280744", "10,650.2", [
        "Replaced engine oil pressure switch with new, zero-time switch P/N 83278, per AD 2013-11-11 paragraph (g)(6) "
        "and Cessna Service Bulletin SB 07-79-01.",
        "Removed switch P/N 83278 at 3,000.4 hours since installation. Ground run: oil pressure indication normal, "
        "no leaks.",
        "Torque per maintenance manual. Wrench T-014.",
    ], "Y. Alaoui", "A. Tazi"),
    ("WO-2024-0071", "14 February 2024", "CN-TFA", "Cessna 172S", "172S9488", "7,210.4", [
        "Replaced engine oil pressure switch with new, zero-time switch P/N 83278, per AD 2013-11-11 paragraph (g)(6).",
        "Ground run: oil pressure indication normal, no leaks. AD compliance record updated.",
        "Torque per maintenance manual. Wrench T-022.",
    ], "S. Idrissi", "A. Tazi"),
    ("WO-2026-0388", "6 October 2026", "CN-TFC", "Cessna 172S", "172S11290", "6,184.9", [
        "Customer request: inspect muffler for cracks after a carbon monoxide detector alert in flight.",
        "Muffler AWI P/N A1754001-23, S/N 34729. Not among the serial numbers listed in AD 2018-02-04, figure 1. "
        "Inspected the weld of the muffler body to the cuff following the method of AD 2018-02-04 paragraph (g)(1) "
        "anyway: no cracks, no broken welds.",
        "Replaced the carbon monoxide detector card (expired). Cabin heat checked on ground run.",
    ], "Y. Alaoui", "K. Bennani"),
]


# --- Builders ---------------------------------------------------------------------------------

def reproducible_zip(path: Path) -> None:
    """Office files are zip archives stamped with the build time; fix the stamp."""
    with zipfile.ZipFile(path) as src:
        entries = [(info, src.read(info)) for info in src.infolist()]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as dst:
        for info, data in entries:
            fixed = zipfile.ZipInfo(info.filename, date_time=FIXED.timetuple()[:6])
            fixed.compress_type = zipfile.ZIP_DEFLATED
            dst.writestr(fixed, data)


def workbook(path: Path, title: str, sheets: list[tuple[str, list[str], list[tuple]]]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.title, wb.properties.creator = title, COMPANY
    wb.properties.created = wb.properties.modified = FIXED
    for name, header, rows in sheets:
        ws = wb.create_sheet(name)
        ws.append(header)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in rows:
            ws.append(list(row))
        for column in ws.columns:
            ws.column_dimensions[column[0].column_letter].width = min(60, max(len(str(c.value or "")) for c in column) + 2)
    wb.save(path)
    reproducible_zip(path)


def build_fleet(path: Path) -> None:
    workbook(path, "Tadla Flight Academy fleet register, 26 October 2026", [
        ("Fleet", ["Registration", "Model", "Serial number", "Year", "Airframe total time (h)", "Engine",
                   "Engine serial number", "Engine time since overhaul (h)", "Base", "Notes"], FLEET),
        ("Equipment", ["Registration", "Item", "Manufacturer", "Model / part number", "Serial number", "Software",
                       "Installed", "Notes"], EQUIPMENT),
        ("Time-controlled parts", ["Registration", "Part", "Part number", "Serial number", "Installed (date)",
                                   "Airframe hours at installation", "Limit", "Source of limit"], TIME_CONTROLLED),
    ])


def build_prices(path: Path) -> None:
    workbook(path, "Parts and labour price list 2026 (MAD)", [
        ("Prices 2026", ["Part number", "Description", "Supplier", "Unit cost (MAD)", "Customer price (MAD)",
                         "Lead time (days)", "Notes"], PRICES),
    ])


def build_procedures(path: Path) -> None:
    from docx import Document

    document = Document()
    document.core_properties.title = "Repair station procedures manual, revision 7 (excerpt)"
    document.core_properties.author = COMPANY
    document.core_properties.created = document.core_properties.modified = FIXED
    document.add_heading(f"{COMPANY}: procedures manual, revision 7 (excerpt)", level=1)
    document.add_paragraph("Effective 1 June 2026. Controlled copy: print only for use on the hangar floor.")
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, ("Revision", "Date", "Change")):
        cell.text = text
    for revision in REVISIONS:
        for cell, text in zip(table.add_row().cells, revision):
            cell.text = text
    for heading, paragraphs in PROCEDURES:
        document.add_heading(heading, level=2)
        for text in paragraphs:
            document.add_paragraph(text)
    document.save(path)
    reproducible_zip(path)


def build_training(path: Path) -> None:
    from pptx import Presentation

    deck = Presentation()
    deck.core_properties.title = "Airworthiness directives at Anfa: training for new mechanics"
    deck.core_properties.author = COMPANY
    deck.core_properties.created = deck.core_properties.modified = FIXED
    for title, body, notes in SLIDES:
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
        slide.notes_slide.notes_text_frame.text = notes
    deck.save(path)
    reproducible_zip(path)


def build_work_orders(path: Path) -> None:
    """One work order per page, rendered as an image and degraded like a photocopy: no text layer."""
    from PIL import Image, ImageDraw, ImageFilter, ImageFont
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen.canvas import Canvas

    rng = random.Random(11)
    width, height = 827, 1169  # A4 at 100 dpi
    small, body, big = (ImageFont.load_default(size=s) for s in (14, 17, 24))
    canvas = Canvas(str(path), pagesize=A4, invariant=1)
    canvas.setTitle("Work orders, scanned")
    for number, date, reg, model, serial, hours, lines, mechanic, inspector in WORK_ORDERS:
        img = Image.new("L", (width, height), 250)
        draw = ImageDraw.Draw(img)
        draw.text((70, 60), COMPANY.upper(), fill=30, font=big)
        draw.text((70, 92), "Hangar 3, Casablanca  |  Work order", fill=60, font=small)
        draw.text((560, 60), number, fill=30, font=big)
        fields = [("Date", date), ("Registration", reg), ("Aircraft", model), ("Serial number", serial),
                  ("Airframe total time at completion (h)", hours)]
        y = 150
        for label, value in fields:
            draw.rectangle([70, y, 757, y + 34], outline=90, width=1)
            draw.text((80, y + 9), label, fill=70, font=small)
            draw.text((400, y + 7), value, fill=25, font=body)
            y += 34
        y += 30
        draw.text((70, y), "Work performed", fill=40, font=body)
        y += 34
        for text in lines:
            line = ""
            for word in text.split():
                trial = f"{line} {word}".strip()
                if draw.textlength(trial, font=body) > width - 160:
                    draw.text((80, y), line, fill=35, font=body)
                    y, line = y + 25, word
                else:
                    line = trial
            draw.text((80, y), line, fill=35, font=body)
            y += 40
        y = max(y + 30, 820)
        for label, name in (("Performed by", mechanic), ("Approved for return to service", inspector)):
            draw.text((80, y), f"{label}: {name}", fill=35, font=body)
            draw.line([(460, y + 22), (740, y + 22)], fill=80, width=1)
            draw.text((470, y - 2), name.split()[-1].lower(), fill=20, font=big)  # the signature
            y += 60
        draw.ellipse([600, 1000, 740, 1080], outline=110, width=3)  # the stamp
        draw.text((622, 1030), "QUALITY", fill=110, font=small)
        img = img.rotate(rng.uniform(-1.1, 1.1), expand=False, fillcolor=245)
        img = img.filter(ImageFilter.GaussianBlur(0.7))
        noise = Image.effect_noise((width, height), 24).convert("L")
        img = Image.blend(img, noise, 0.1)
        buf = io.BytesIO()
        img.save(buf, "JPEG", quality=32)
        buf.seek(0)
        canvas.drawImage(ImageReader(buf), 0, 0, width=A4[0], height=A4[1])
        canvas.showPage()
    canvas.save()


# --- Records from the maintenance history -------------------------------------------------------

def _pdf(path: Path, story: list, title: str, landscape_page: bool = False) -> None:
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate

    size = landscape(A4) if landscape_page else A4
    doc = SimpleDocTemplate(str(path), pagesize=size, title=title, author=COMPANY, invariant=1,
                            leftMargin=18 * mm, rightMargin=18 * mm, topMargin=16 * mm, bottomMargin=16 * mm)
    doc.build(story)


def _styles():
    from reportlab.lib.styles import getSampleStyleSheet

    styles = getSampleStyleSheet()
    return styles["Title"], styles["Heading3"], styles["BodyText"]


def _grid(rows: list[list], widths: list[float]):
    from reportlab.lib import colors
    from reportlab.platypus import Table, TableStyle

    table = Table(rows, colWidths=widths, repeatRows=1)
    table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.grey), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                               ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e8e8e8")),
                               ("FONTSIZE", (0, 0), (-1, -1), 8.5)]))
    return table


def build_work_order(path: Path, order, history) -> None:
    """A work order as the maintenance system prints it: a text layer, labelled fields, a parts table."""
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, Spacer

    title, heading, body = _styles()
    model = {reg: model for reg, model, *_ in FLEET}[order.reg]
    serial = {reg: serial for reg, _, serial, *_ in FLEET}[order.reg]
    story = [Paragraph(f"{COMPANY}: work order {order.number}", title),
             _grid([["Field", "Value"],
                    ["Date completed", order.day.strftime("%d %B %Y").lstrip("0")],
                    ["Registration", order.reg], ["Aircraft", f"Cessna {model}"], ["Serial number", serial],
                    ["Airframe total time at completion (h)", f"{order.hours:,.1f}"],
                    ["Station", history.STATION.get(order.reg, "Casablanca")]], [70 * mm, 100 * mm]),
             Paragraph("Request", heading), Paragraph(order.request, body), Paragraph("Work performed", heading)]
    story += [Paragraph(f"{i}. {line}", body) for i, line in enumerate(order.lines, 1)]
    if order.parts:
        story += [Paragraph("Parts used", heading),
                  _grid([["Part number", "Description", "Quantity", "Serial number"],
                         *[[pn, desc, str(qty), sn] for pn, desc, qty, sn in order.parts]],
                        [40 * mm, 70 * mm, 25 * mm, 35 * mm])]
    licence = history.INSPECTORS[order.inspector]
    story += [Spacer(1, 6 * mm),
              Paragraph("Certification: the work above was performed in accordance with 14 CFR Part 43 and the "
                        "procedures manual. The aircraft is airworthy for the work performed.", body),
              Paragraph(f"Performed by: {order.mechanic}", body),
              Paragraph(f"Return to service: {order.inspector}, licence {licence}", body)]
    _pdf(path, story, f"Work order {order.number}")


def build_logbook(path: Path, reg: str, orders, history) -> None:
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph

    title, _, body = _styles()
    rows = [["Date", "Airframe hours", "Work order", "Entry", "Signed"]]
    for order in orders:
        rows.append([order.day.isoformat(), f"{order.hours:,.1f}", order.number, Paragraph(history.logbook_entry(order), body),
                     Paragraph(f"{order.inspector}, {history.INSPECTORS[order.inspector]}", body)])
    story = [Paragraph(f"{reg} airframe logbook: entries from 1 January 2024 to 26 October 2026 (printed extract)", title),
             _grid(rows, [24 * mm, 30 * mm, 30 * mm, 150 * mm, 35 * mm])]
    _pdf(path, story, f"{reg} airframe logbook extract", landscape_page=True)


def build_ad_record(path: Path, reg: str, history) -> None:
    columns = ["AD number", "Effective", "Subject", "Applies", "Notes", "Method of compliance", "Date",
               "Airframe hours", "Recurring", "Next due", "Signed"]
    keys = ["ad", "effective", "subject", "applies", "notes", "method", "date", "hours", "recurring", "next_due", "signed"]
    rows = [tuple(row[k] for k in keys) for row in history.ad_record(reg)]
    workbook(path, f"{reg} AD compliance record, 26 October 2026", [("AD compliance record", columns, rows)])


def build_records(catalog: list[dict]) -> None:
    import anfa_history as history

    orders = history.work_orders()
    folders = {name: OUT / name for name in ("work-orders", "logbooks", "ad-records")}
    for folder in folders.values():
        folder.mkdir(parents=True, exist_ok=True)
        for old in folder.iterdir():
            old.unlink()
    kinds = {"100-hour": "100-hour inspection", "annual": "annual inspection", "oil pressure switch": "oil pressure switch replacement",
             "engine overhaul": "engine overhaul", "fadec battery": "FADEC backup battery replacement",
             "transponder": "altimeter, static system, and transponder tests", "unscheduled": "repair"}
    for order in orders:
        if order.scanned:
            continue
        name = f"work-orders/{order.number}.pdf"
        build_work_order(OUT / name, order, history)
        catalog.append({"source_id": order.number, "file": f"anfa/{name}", "meta": {
            "publisher": COMPANY, "title": f"Work order {order.number}: {order.reg}, {kinds[order.kind]}",
            "doc_type": "work order", "aircraft": order.reg, "date": order.day.isoformat(), "access": "staff"}})
    for reg in history.ANCHORS:
        name = f"logbooks/logbook-{reg}.pdf"
        build_logbook(OUT / name, reg, [o for o in orders if o.reg == reg], history)
        catalog.append({"source_id": f"LOG-{reg}", "file": f"anfa/{name}", "meta": {
            "publisher": COMPANY, "title": f"{reg} airframe logbook extract, 2024 to 2026", "doc_type": "logbook",
            "aircraft": reg, "as_of": "2026-10-26", "access": "staff"}})
        name = f"ad-records/ad-record-{reg}.xlsx"
        build_ad_record(OUT / name, reg, history)
        catalog.append({"source_id": f"ADREC-{reg}", "file": f"anfa/{name}", "meta": {
            "publisher": COMPANY, "title": f"{reg} AD compliance record, 26 October 2026", "doc_type": "AD compliance record",
            "aircraft": reg, "as_of": "2026-10-26", "access": "staff"}})
    print(f"records        -> {sum(not o.scanned for o in orders)} work orders, {len(history.ANCHORS)} logbooks, "
          f"{len(history.ANCHORS)} AD compliance records")


DOCUMENTS = [
    ("FLEET-2026-10", "fleet-register-2026-10.xlsx", build_fleet,
     {"title": "Tadla Flight Academy fleet register, 26 October 2026", "doc_type": "fleet register", "as_of": "2026-10-26", "access": "staff"}),
    ("PROC-R7", "procedures-manual-r7.docx", build_procedures,
     {"title": "Procedures manual, revision 7 (excerpt)", "doc_type": "procedures manual", "revision": "7", "effective": "2026-06-01", "access": "employees"}),
    ("TRAIN-AD", "training-airworthiness-directives.pptx", build_training,
     {"title": "Airworthiness directives at Anfa: training deck for new mechanics", "doc_type": "training", "access": "employees"}),
    ("WO-SCAN", "work-orders-scanned.pdf", build_work_orders,
     {"title": "Work orders WO-2021-0544, WO-2024-0071, WO-2026-0388 (scanned)", "doc_type": "work orders", "access": "staff"}),
    ("PRICES-2026", "price-list-2026.xlsx", build_prices,
     {"title": "Parts and labour price list 2026", "doc_type": "price list", "access": "management"}),
]


def main() -> None:
    from reportlab import rl_config

    rl_config.invariant = 1  # no build time or random ids in the PDFs
    OUT.mkdir(parents=True, exist_ok=True)
    catalog_path = SOURCE / "catalog.json"
    catalog = json.loads(catalog_path.read_text(encoding="utf-8"))["documents"] if catalog_path.exists() else []
    catalog = [d for d in catalog if not d["file"].startswith("anfa/")]
    for source_id, name, build, meta in DOCUMENTS:
        path = OUT / name
        build(path)
        catalog.append({"source_id": source_id, "file": f"anfa/{name}", "meta": {"publisher": COMPANY, **meta}})
        print(f"{source_id:14} -> {path.relative_to(ROOT).as_posix()} ({path.stat().st_size // 1024} KB)")
    build_records(catalog)
    catalog.sort(key=lambda d: d["file"])
    catalog_path.write_bytes((json.dumps({"documents": catalog}, indent=2) + "\n").encode("utf-8"))


if __name__ == "__main__":
    main()
